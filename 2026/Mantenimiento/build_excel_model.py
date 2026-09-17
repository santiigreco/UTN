# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

def create_model():
    wb = openpyxl.Workbook()

    # Setup Sheet
    ws = wb.active
    ws.title = 'Factibilidad Economica'
    ws.views.sheetView[0].showGridLines = True

    # Paleta de Colores Corporativa UTN
    NAVY = '1B365D'
    BLUE_LIGHT = 'E8EEF5'
    BLUE_HEADER = '2E5B88'
    GRAY_BG = 'F8FAFC'
    GRAY_BORDER = 'CBD5E0'
    WHITE = 'FFFFFF'
    GREEN_KPI = '00875A'
    GREEN_BG = 'E3FCEF'

    font_title = Font(name='Calibri', size=15, bold=True, color=NAVY)
    font_subtitle = Font(name='Calibri', size=10, italic=True, color='4A5568')
    font_sec_hdr = Font(name='Calibri', size=11, bold=True, color=WHITE)
    font_tbl_hdr = Font(name='Calibri', size=10, bold=True, color=WHITE)
    font_bold = Font(name='Calibri', size=10, bold=True, color='1A202C')
    font_regular = Font(name='Calibri', size=10, color='2D3748')
    font_note = Font(name='Calibri', size=9, italic=True, color='718096')

    font_kpi_val = Font(name='Calibri', size=13, bold=True, color=NAVY)
    font_kpi_val_green = Font(name='Calibri', size=13, bold=True, color=GREEN_KPI)

    fill_navy = PatternFill(start_color=NAVY, end_color=NAVY, fill_type='solid')
    fill_blue_hdr = PatternFill(start_color=BLUE_HEADER, end_color=BLUE_HEADER, fill_type='solid')
    fill_subtotal = PatternFill(start_color=BLUE_LIGHT, end_color=BLUE_LIGHT, fill_type='solid')
    fill_gray = PatternFill(start_color=GRAY_BG, end_color=GRAY_BG, fill_type='solid')
    fill_kpi_green = PatternFill(start_color=GREEN_BG, end_color=GREEN_BG, fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color=GRAY_BORDER),
        right=Side(style='thin', color=GRAY_BORDER),
        top=Side(style='thin', color=GRAY_BORDER),
        bottom=Side(style='thin', color=GRAY_BORDER)
    )

    double_bottom = Border(
        left=Side(style='thin', color=GRAY_BORDER),
        right=Side(style='thin', color=GRAY_BORDER),
        top=Side(style='thin', color=GRAY_BORDER),
        bottom=Side(style='double', color=NAVY)
    )

    # Header Superior
    ws['A2'] = 'UNIVERSIDAD TECNOLÓGICA NACIONAL - FRBA | DEPARTAMENTO DE INGENIERÍA INDUSTRIAL'
    ws['A2'].font = font_title
    ws['A3'] = 'Cátedra: Mantenimiento (Curso I-5052) | Grupo Nº 5 | Empresa: UnionBat S.A.'
    ws['A3'].font = font_subtitle
    ws['A4'] = 'MÓDULO 2 - PUNTO 1.3: ANÁLISIS DE FACTIBILIDAD TÉCNICO-ECONÓMICO (CASH FLOW & TIR)'
    ws['A4'].font = Font(name='Calibri', size=11, bold=True, color=BLUE_HEADER)

    # Formatos
    FMT_CURR = '"$"#,##0'
    FMT_CURR_NEG = '"$"#,##0;[Red]("$"#,##0);"-"'
    FMT_PCT = '0.0%'

    # ==========================================
    # 1. PARAMETROS BASE (DEL TP)
    # ==========================================
    ws['A6'] = '1. PARÁMETROS BASE DE OPERACIÓN Y COSTOS (UNIONBAT S.A.)'
    ws.merge_cells('A6:E6')
    ws['A6'].font = font_sec_hdr
    ws['A6'].fill = fill_navy
    ws['A6'].alignment = Alignment(vertical='center', indent=1)

    params_data = [
        ('Facturación Anual de la Planta (605.000 bat/año @ USD 100)', 60500000, FMT_CURR, 'Punto 1.4 del TP ($84.700 M ARS)'),
        ('Presupuesto Anual de Mantenimiento (5% s/Ventas)', 3025000, FMT_CURR, 'Puntos 1.10 y 2.2 ($4.235 M ARS)'),
        ('Presupuesto de Materiales y Repuestos (30% s/Mantenimiento)', 907500, FMT_CURR, 'Punto 2.2 Tabla 4 ($1.271 M ARS)'),
        ('Costos Indirectos de Fabricación CIF (12% s/Ventas)', 7260000, FMT_CURR, 'Punto 1.10 Tabla 2 ($10.164 M ARS)'),
        ('Margen Operativo de la Empresa (4% s/Ventas)', 2420000, FMT_CURR, 'Punto 1.10 Tabla 2 ($3.388 M ARS)'),
        ('Tasa de Descuento / Tasa de Corte (WACC en USD)', 0.12, FMT_PCT, 'Tasa de referencia de costo de capital'),
        ('Horizonte de Evaluación del Proyecto', 5, '0 "Años"', 'Vida útil proyectada de las mejoras')
    ]

    for i, (lbl, val, fmt, ref) in enumerate(params_data, start=7):
        ws[f'A{i}'] = lbl
        ws[f'B{i}'] = val
        ws[f'B{i}'].number_format = fmt
        ws[f'B{i}'].alignment = Alignment(horizontal='right')
        ws[f'C{i}'] = ref
        ws[f'A{i}'].font = font_regular
        ws[f'B{i}'].font = font_bold
        ws[f'C{i}'].font = font_note
        for c in ['A', 'B', 'C']:
            ws[f'{c}{i}'].border = thin_border

    # ==========================================
    # 2. INVERSION INICIAL (CAPEX) - TABLAS 5 A 9
    # ==========================================
    ws['A15'] = '2. PRESUPUESTO DE INVERSIÓN INICIAL DE REINGENIERÍA (CAPEX - AÑO 0)'
    ws.merge_cells('A15:E15')
    ws['A15'].font = font_sec_hdr
    ws['A15'].fill = fill_navy
    ws['A15'].alignment = Alignment(vertical='center', indent=1)

    ws['A16'] = 'Eje de Reingeniería / Componente'
    ws['B16'] = 'Monto [USD]'
    ws['C16'] = 'Participación %'
    ws['D16'] = 'Referencia en Documento del TP'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}16'].font = font_tbl_hdr
        ws[f'{c}16'].fill = fill_blue_hdr
        ws[f'{c}16'].alignment = Alignment(horizontal='center' if c in ['B', 'C'] else 'left')
        ws[f'{c}16'].border = thin_border

    capex_items = [
        ('Eje 1: Bateas y Rectificadores Digatron (Intercambiador, mordazas EPDM, pirómetros SCADA)', 38000, 'Punto 1.2.1 Tabla 5'),
        ('Eje 2: Sistema de Extracción y Ventilación (Scrubber PP/PEAD, ductos PRFV, VDF, motores IP66)', 45000, 'Punto 1.2.1 Tabla 6'),
        ('Eje 3: Lavadoras Industriales MAC (Bombas magnéticas PVDF, filtro doble, cabina acústica)', 22000, 'Punto 1.2.1 Tabla 7'),
        ('Eje 4: Mecanización y Transporte KUKA (Motorreductores norma, garras PU, aire seco sensores)', 15000, 'Punto 1.2.1 Tabla 8'),
    ]

    for i, (eje, val, ref) in enumerate(capex_items, start=17):
        ws[f'A{i}'] = eje
        ws[f'B{i}'] = val
        ws[f'B{i}'].number_format = FMT_CURR
        ws[f'C{i}'] = f'=B{i}/$B$23'
        ws[f'C{i}'].number_format = FMT_PCT
        ws[f'D{i}'] = ref
        ws[f'A{i}'].font = font_regular
        ws[f'B{i}'].font = font_regular
        ws[f'C{i}'].font = font_regular
        ws[f'D{i}'].font = font_note
        ws[f'A{i}'].alignment = Alignment(horizontal='left')
        ws[f'B{i}'].alignment = Alignment(horizontal='right')
        ws[f'C{i}'].alignment = Alignment(horizontal='center')
        for c in ['A', 'B', 'C', 'D']:
            ws[f'{c}{i}'].border = thin_border

    # Subtotal equipos
    ws['A21'] = 'Subtotal Equipamiento y Materiales'
    ws['B21'] = '=SUM(B17:B20)'
    ws['B21'].number_format = FMT_CURR
    ws['C21'] = '=B21/$B$23'
    ws['C21'].number_format = FMT_PCT
    ws['D21'] = 'Suma de módulos técnicos'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}21'].font = font_bold
        ws[f'{c}21'].fill = fill_gray
        ws[f'{c}21'].border = thin_border

    # Montaje e imprevistos
    ws['A22'] = 'Montaje, Integración SCADA e Imprevistos de Obra (16,7% s/Equipos)'
    ws['B22'] = 20000
    ws['B22'].number_format = FMT_CURR
    ws['C22'] = '=B22/$B$23'
    ws['C22'].number_format = FMT_PCT
    ws['D22'] = 'Cronograma Tabla 9 (Desembolso $140.000 total)'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}22'].font = font_regular
        ws[f'{c}22'].border = thin_border

    # Total CAPEX
    ws['A23'] = 'TOTAL INVERSIÓN INICIAL DE REINGENIERÍA (CAPEX t=0)'
    ws['B23'] = '=B21+B22'
    ws['B23'].number_format = FMT_CURR
    ws['C23'] = '=B23/$B$23'
    ws['C23'].number_format = FMT_PCT
    ws['D23'] = 'Total desembolsos programados semanas 1-12'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}23'].font = font_bold
        ws[f'{c}23'].fill = fill_subtotal
        ws[f'{c}23'].border = double_bottom

    # ==========================================
    # 3. BENEFICIOS Y AHORROS ANUALES (CASH INFLOWS)
    # ==========================================
    ws['A25'] = '3. AHORROS Y BENEFICIOS OPERATIVOS ANUALES ESTIMADOS (INGRESOS DEL PROYECTO)'
    ws.merge_cells('A25:E25')
    ws['A25'].font = font_sec_hdr
    ws['A25'].fill = fill_navy
    ws['A25'].alignment = Alignment(vertical='center', indent=1)

    ws['A26'] = 'Concepto de Ahorro / Beneficio Cuantificado'
    ws['B26'] = 'Ahorro Anual [USD]'
    ws['C26'] = 'Participación %'
    ws['D26'] = 'Fundamento Técnico en TP'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}26'].font = font_tbl_hdr
        ws[f'{c}26'].fill = fill_blue_hdr
        ws[f'{c}26'].alignment = Alignment(horizontal='center' if c in ['B', 'C'] else 'left')
        ws[f'{c}26'].border = thin_border

    savings_items = [
        ('Aumento de Capacidad (+12%) y eliminación de paradas térmicas (85°C) en cuello de botella (Digatron)', 65000, 'Reducción de ciclo y paradas en sector activación (P244 y Tabla 9)'),
        ('Ahorro en Mantenimiento Correctivo y Repuestos (Bombas PVDF, Scrubber, rodamientos)', 45000, 'Mitigación de desgaste sobre ppto de repuestos de USD 907.500 (Tabla 4)'),
        ('Eficiencia Energética por modulación VDF y motores directos IP66 en extractores', 18000, 'Ahorro eléctrico en CIF de energía (12% costos - P189 y P255)'),
        ('Reducción de MTTR (<15 min) y eliminación de scrap por rotura en celda robótica KUKA', 12000, 'Estandarización de motores y garras PU vulcanizadas (P277 y P281)'),
    ]

    for i, (con, val, ref) in enumerate(savings_items, start=27):
        ws[f'A{i}'] = con
        ws[f'B{i}'] = val
        ws[f'B{i}'].number_format = FMT_CURR
        ws[f'C{i}'] = f'=B{i}/$B$31'
        ws[f'C{i}'].number_format = FMT_PCT
        ws[f'D{i}'] = ref
        ws[f'A{i}'].font = font_regular
        ws[f'B{i}'].font = font_regular
        ws[f'C{i}'].font = font_regular
        ws[f'D{i}'].font = font_note
        ws[f'A{i}'].alignment = Alignment(horizontal='left')
        ws[f'B{i}'].alignment = Alignment(horizontal='right')
        ws[f'C{i}'].alignment = Alignment(horizontal='center')
        for c in ['A', 'B', 'C', 'D']:
            ws[f'{c}27'].border = thin_border
            ws[f'{c}{i}'].border = thin_border

    # Total Beneficios Brutos
    ws['A31'] = 'TOTAL BENEFICIOS OPERATIVOS BRUTOS ANUALES'
    ws['B31'] = '=SUM(B27:B30)'
    ws['B31'].number_format = FMT_CURR
    ws['C31'] = '=B31/$B$31'
    ws['C31'].number_format = FMT_PCT
    ws['D31'] = 'Suma de ahorros anuales proyectados'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}31'].font = font_bold
        ws[f'{c}31'].fill = fill_gray
        ws[f'{c}31'].border = thin_border

    # OPEX Incremental
    ws['A32'] = '(-) Costo Operativo de Mantenimiento Incremental (OPEX de nuevas tecnologías)'
    ws['B32'] = 10000
    ws['B32'].number_format = FMT_CURR
    ws['C32'] = '=B32/$B$31'
    ws['C32'].number_format = FMT_PCT
    ws['D32'] = 'Insumos químicos Scrubber, filtros dobles y aire seco'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}32'].font = font_regular
        ws[f'{c}32'].border = thin_border

    # Beneficio Neto Anual
    ws['A33'] = 'BENEFICIO NETO OPERATIVO ANUAL (CASH INFLOW NETO)'
    ws['B33'] = '=B31-B32'
    ws['B33'].number_format = FMT_CURR
    ws['C33'] = '=B33/$B$31'
    ws['C33'].number_format = FMT_PCT
    ws['D33'] = 'Flujo de fondos operativo anual generado'
    for c in ['A', 'B', 'C', 'D']:
        ws[f'{c}33'].font = font_bold
        ws[f'{c}33'].fill = fill_subtotal
        ws[f'{c}33'].border = double_bottom

    # ==========================================
    # 4. TABLA DE CASH FLOW PROYECTADO (AÑOS 0 A 5)
    # ==========================================
    ws['A36'] = '4. MATRIZ DE FLUJO DE FONDOS PROYECTADO (CASH FLOW EN USD)'
    ws.merge_cells('A36:G36')
    ws['A36'].font = font_sec_hdr
    ws['A36'].fill = fill_navy
    ws['A36'].alignment = Alignment(vertical='center', indent=1)

    cf_headers = ['Concepto / Período', 'Año 0', 'Año 1', 'Año 2', 'Año 3', 'Año 4', 'Año 5']
    for col_idx, h in enumerate(cf_headers, start=1):
        col_letter = get_column_letter(col_idx)
        ws[f'{col_letter}37'] = h
        ws[f'{col_letter}37'].font = font_tbl_hdr
        ws[f'{col_letter}37'].fill = fill_blue_hdr
        ws[f'{col_letter}37'].alignment = Alignment(horizontal='center' if col_idx > 1 else 'left')
        ws[f'{col_letter}37'].border = thin_border

    # Fila: Inversión Inicial CAPEX
    ws['A38'] = '(-) Inversión Inicial (CAPEX)'
    ws['B38'] = '=-B23'
    ws['B38'].number_format = FMT_CURR_NEG
    for c_idx in range(3, 8):
        c_let = get_column_letter(c_idx)
        ws[f'{c_let}38'] = 0
        ws[f'{c_let}38'].number_format = FMT_CURR_NEG

    # Fila: Beneficios Brutos
    ws['A39'] = '(+) Ahorros y Beneficios Operativos'
    ws['B39'] = 0
    ws['B39'].number_format = FMT_CURR_NEG
    for c_idx in range(3, 8):
        c_let = get_column_letter(c_idx)
        ws[f'{c_let}39'] = '=$B$31'
        ws[f'{c_let}39'].number_format = FMT_CURR

    # Fila: OPEX
    ws['A40'] = '(-) Costos de Mantenimiento Incremental (OPEX)'
    ws['B40'] = 0
    ws['B40'].number_format = FMT_CURR_NEG
    for c_idx in range(3, 8):
        c_let = get_column_letter(c_idx)
        ws[f'{c_let}40'] = '=-$B$32'
        ws[f'{c_let}40'].number_format = FMT_CURR_NEG

    # Fila: Flujo de Fondos Neto (FNC)
    ws['A41'] = 'FLUJO NETO DE CAJA (FNC)'
    for c_idx in range(2, 8):
        c_let = get_column_letter(c_idx)
        ws[f'{c_let}41'] = f'=SUM({c_let}38:{c_let}40)'
        ws[f'{c_let}41'].number_format = FMT_CURR_NEG
        ws[f'{c_let}41'].font = font_bold
        ws[f'{c_let}41'].fill = fill_subtotal
        ws[f'{c_let}41'].border = thin_border
    ws['A41'].font = font_bold
    ws['A41'].fill = fill_subtotal
    ws['A41'].border = thin_border

    # Fila: Flujo Acumulado
    ws['A42'] = 'Flujo de Fondos Acumulado'
    ws['B42'] = '=B41'
    ws['B42'].number_format = FMT_CURR_NEG
    ws['B42'].font = font_bold
    ws['B42'].border = thin_border
    for c_idx in range(3, 8):
        prev_let = get_column_letter(c_idx - 1)
        curr_let = get_column_letter(c_idx)
        ws[f'{curr_let}42'] = f'={prev_let}42+{curr_let}41'
        ws[f'{curr_let}42'].number_format = FMT_CURR_NEG
        ws[f'{curr_let}42'].font = font_bold
        ws[f'{curr_let}42'].border = thin_border
    ws['A42'].font = font_bold
    ws['A42'].border = thin_border

    # Fila: Factor de Descuento
    ws['A43'] = 'Factor de Descuento (1 / (1+WACC)^t)'
    ws['B43'] = 1.0
    ws['B43'].number_format = '0.0000'
    for c_idx in range(3, 8):
        t_val = c_idx - 2
        c_let = get_column_letter(c_idx)
        ws[f'{c_let}43'] = f'=1/(1+$B$12)^{t_val}'
        ws[f'{c_let}43'].number_format = '0.0000'
    for c_idx in range(1, 8):
        ws[f'{get_column_letter(c_idx)}43'].font = font_note
        ws[f'{get_column_letter(c_idx)}43'].border = thin_border

    # Fila: Flujo Descontado
    ws['A44'] = 'Flujo Neto Descontado (Valor Presente)'
    ws['B44'] = '=B41*B43'
    ws['B44'].number_format = FMT_CURR_NEG
    for c_idx in range(3, 8):
        c_let = get_column_letter(c_idx)
        ws[f'{c_let}44'] = f'={c_let}41*{c_let}43'
        ws[f'{c_let}44'].number_format = FMT_CURR_NEG
    for c_idx in range(1, 8):
        ws[f'{get_column_letter(c_idx)}44'].font = font_regular
        ws[f'{get_column_letter(c_idx)}44'].border = thin_border

    # Fila: Flujo Descontado Acumulado
    ws['A45'] = 'Flujo Descontado Acumulado'
    ws['B45'] = '=B44'
    ws['B45'].number_format = FMT_CURR_NEG
    for c_idx in range(3, 8):
        prev_let = get_column_letter(c_idx - 1)
        curr_let = get_column_letter(c_idx)
        ws[f'{curr_let}45'] = f'={prev_let}45+{curr_let}44'
        ws[f'{curr_let}45'].number_format = FMT_CURR_NEG
    for c_idx in range(1, 8):
        ws[f'{get_column_letter(c_idx)}45'].font = font_bold
        ws[f'{get_column_letter(c_idx)}45'].border = double_bottom

    for r in range(38, 41):
        ws[f'A{r}'].font = font_regular
        ws[f'A{r}'].border = thin_border
        for c_idx in range(2, 8):
            ws[f'{get_column_letter(c_idx)}{r}'].font = font_regular
            ws[f'{get_column_letter(c_idx)}{r}'].border = thin_border

    # ==========================================
    # 5. TABLERO DE INDICADORES FINANCIEROS (KPIS)
    # ==========================================
    ws['A48'] = '5. TABLERO DE INDICADORES DE RENTABILIDAD Y FACTIBILIDAD ECONÓMICA'
    ws.merge_cells('A48:G48')
    ws['A48'].font = font_sec_hdr
    ws['A48'].fill = fill_navy
    ws['A48'].alignment = Alignment(vertical='center', indent=1)

    kpis = [
        ('TASA INTERNA DE RETORNO (TIR)', '=IRR(B41:G41)', FMT_PCT, 'Supera ampliamente el WACC (12,0%)', font_kpi_val_green),
        ('VALOR ACTUAL NETO (VAN @ 12%)', '=NPV(B12,C41:G41)+B41', FMT_CURR, 'Riqueza neta generada para UnionBat', font_kpi_val_green),
        ('PERÍODO DE RECUPERO SIMPLE (PAYBACK)', '=ABS(B41)/C41', '0.00 "Años" (aprox. 13 meses)', 'Recupero veloz de la inversión inicial', font_kpi_val),
        ('PERÍODO DE RECUPERO DESCONTADO', '=(ABS(B41)-C44)/D44 + 1', '0.00 "Años" (aprox. 15 meses)', 'Considerando el costo del capital', font_kpi_val),
        ('RELACIÓN BENEFICIO / COSTO (B/C)', '=(NPV(B12,C41:G41))/ABS(B41)', '0.00', 'Por cada USD 1 invertido se obtienen USD 3,35', font_kpi_val),
        ('DICTAMEN DE FACTIBILIDAD', 'PROYECTO ALTAMENTE FACTIBLE Y RENTABLE', None, 'Se recomienda aprobación y ejecución inmediata', Font(name='Calibri', size=11, bold=True, color=GREEN_KPI))
    ]

    for idx, (k_lbl, k_val, k_fmt, k_comm, k_fnt) in enumerate(kpis, start=49):
        ws[f'A{idx}'] = k_lbl
        ws[f'C{idx}'] = k_val
        if k_fmt:
            ws[f'C{idx}'].number_format = k_fmt
        ws[f'D{idx}'] = k_comm
        ws.merge_cells(f'A{idx}:B{idx}')
        ws.merge_cells(f'D{idx}:G{idx}')
        
        ws[f'A{idx}'].font = font_bold
        ws[f'C{idx}'].font = k_fnt
        ws[f'D{idx}'].font = font_subtitle
        ws[f'C{idx}'].alignment = Alignment(horizontal='center')
        
        for c_idx in range(1, 8):
            ws[f'{get_column_letter(c_idx)}{idx}'].border = thin_border
            if idx == 54:
                ws[f'{get_column_letter(c_idx)}{idx}'].fill = fill_kpi_green

    # ==========================================
    # 6. GRAFICOS DINAMICOS (INTEGRADOS AL COSTADO)
    # ==========================================

    # GRAFICO 1: Flujo de Fondos Neto por Periodo (Barras)
    chart1 = BarChart()
    chart1.type = 'col'
    chart1.style = 10
    chart1.title = 'Flujo de Fondos Neto Anual (Cash Flow t=0 a t=5)'
    chart1.y_axis.title = 'Monto [USD]'
    chart1.x_axis.title = 'Período'
    chart1.width = 16
    chart1.height = 10

    data1 = Reference(ws, min_col=2, min_row=41, max_col=7, max_row=41)
    cats1 = Reference(ws, min_col=2, min_row=37, max_col=7, max_row=37)
    chart1.add_data(data1, from_rows=True, titles_from_data=False)
    chart1.set_categories(cats1)
    chart1.legend = None
    ws.add_chart(chart1, 'I6')

    # GRAFICO 2: Flujo Acumulado y Payback (Lineas)
    chart2 = LineChart()
    chart2.title = 'Evolución del Cash Flow Acumulado y Punto de Equilibrio (Payback)'
    chart2.style = 13
    chart2.y_axis.title = 'Flujo Acumulado [USD]'
    chart2.x_axis.title = 'Período'
    chart2.width = 16
    chart2.height = 10

    data2 = Reference(ws, min_col=2, min_row=42, max_col=7, max_row=42)
    chart2.add_data(data2, from_rows=True, titles_from_data=False)
    chart2.set_categories(cats1)
    chart2.legend = None
    ws.add_chart(chart2, 'I22')

    # GRAFICO 3: Distribucion de CAPEX por Eje (Torta)
    chart3 = PieChart()
    chart3.title = 'Distribución de la Inversión Inicial (CAPEX) por Eje'
    chart3.width = 16
    chart3.height = 10

    data3 = Reference(ws, min_col=2, min_row=17, max_row=20)
    cats3 = Reference(ws, min_col=1, min_row=17, max_row=20)
    chart3.add_data(data3, titles_from_data=False)
    chart3.set_categories(cats3)
    ws.add_chart(chart3, 'I38')

    # Ajuste de ancho de columnas
    col_widths = {'A': 52, 'B': 22, 'C': 22, 'D': 42, 'E': 14, 'F': 14, 'G': 14, 'H': 3, 'I': 18}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    out_path = r'd:\Descargas\UTN\Repo-UTN\2026\Mantenimiento\TP\Factibilidad_Reingenieria_Grupo5.xlsx'
    wb.save(out_path)
    print(f'Successfully saved to {out_path}')

if __name__ == '__main__':
    create_model()
